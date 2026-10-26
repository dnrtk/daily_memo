#!/usr/bin/env python3
"""memo-index のテスト用フィクスチャ(docx / xlsx / pptx / pdf)を生成する。

必要: python-docx, openpyxl, python-pptx, LibreOffice(PDF 変換用)
    python3 scripts/make-fixtures.py
"""
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

OUT = Path(__file__).resolve().parent.parent / "crates/memo-index/tests/fixtures"
OUT.mkdir(parents=True, exist_ok=True)

# --- docx ---------------------------------------------------------------------
doc = Document()
doc.add_heading("オンプレLLM導入提案書", level=0)
doc.add_heading("背景", level=1)
doc.add_paragraph("顧客はソースコードを社外に出せないため、クラウドの生成AIを利用できない。")
doc.add_heading("構成案", level=1)
doc.add_paragraph("GPUサーバー上で vLLM を稼働させ、社内ネットワークからのみアクセスさせる。")
table = doc.add_table(rows=2, cols=2)
table.cell(0, 0).text = "項目"
table.cell(0, 1).text = "金額"
table.cell(1, 0).text = "GPUサーバー"
table.cell(1, 1).text = "450万円"
doc.add_heading("見積もり", level=2)
doc.add_paragraph("初期費用の概算は 680 万円 & 保守費は年額 120 万円。")
doc.save(OUT / "proposal.docx")

# --- xlsx ---------------------------------------------------------------------
wb = Workbook()
ws = wb.active
ws.title = "見積"
ws.append(["品目", "数量", "単価"])
ws.append(["RTX 6000 Ada", 2, 1200000])
ws.append(["構築作業", 1, 800000])
ws2 = wb.create_sheet("スケジュール")
ws2.append(["フェーズ", "期間"])
ws2.append(["PoC", "2026年11月"])
ws2.append(["本番移行", "2027年1月"])
wb.save(OUT / "estimate.xlsx")

# --- pptx ---------------------------------------------------------------------
prs = Presentation()
s1 = prs.slides.add_slide(prs.slide_layouts[1])
s1.shapes.title.text = "RAG 検索基盤の概要"
s1.placeholders[1].text = "Converter で Office 文書をテキスト化\nハイブリッド検索で取りこぼしを防ぐ"
s1.notes_slide.notes_text_frame.text = "ここで社内事例を口頭で補足する"
s2 = prs.slides.add_slide(prs.slide_layouts[5])
s2.shapes.title.text = "今後の課題"
box = s2.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
box.text_frame.text = "埋め込みモデルの日本語性能を評価する"
prs.save(OUT / "overview.pptx")

# --- pdf(docx を LibreOffice で変換。フォント埋め込みの実際的な PDF になる) ------------
with tempfile.TemporaryDirectory() as tmp:
    src = Path(tmp) / "report.docx"
    d = Document()
    d.add_heading("障害報告書", level=1)
    d.add_paragraph("2026年10月20日、推論サーバーのメモリ不足により応答が停止した。")
    d.add_page_break()
    d.add_heading("再発防止策", level=1)
    d.add_paragraph("KV キャッシュの上限を設定し、監視アラートを追加する。")
    d.save(src)
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, str(src)],
        check=True,
        capture_output=True,
    )
    shutil.copy(Path(tmp) / "report.pdf", OUT / "report.pdf")

print("fixtures written to", OUT)
